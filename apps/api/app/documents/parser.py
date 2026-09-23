import os
from typing import List, Dict, Any
from pypdf import PdfReader
import docx
from pptx import Presentation


def extract_text_and_chunks(file_path: str, mime_type: str, chunk_size: int = 800, chunk_overlap: int = 150) -> List[Dict[str, Any]]:
    chunks = []
    
    # 1. Plain Text
    if "text" in mime_type or file_path.endswith(".txt"):
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            full_text = f.read()
        return _chunk_text(full_text, chunk_size, chunk_overlap)

    # 2. PDF
    elif "pdf" in mime_type or file_path.endswith(".pdf"):
        reader = PdfReader(file_path)
        for page_num, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                page_chunks = _chunk_text(text, chunk_size, chunk_overlap, page_number=page_num)
                chunks.extend(page_chunks)
        return chunks

    # 3. DOCX
    elif "word" in mime_type or file_path.endswith(".docx"):
        doc = docx.Document(file_path)
        full_text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
        return _chunk_text(full_text, chunk_size, chunk_overlap)

    # 4. PPTX
    elif "presentation" in mime_type or file_path.endswith(".pptx") or file_path.endswith(".ppt"):
        prs = Presentation(file_path)
        for slide_num, slide in enumerate(prs.slides, start=1):
            slide_texts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    slide_texts.append(shape.text)
            slide_text = "\n".join(slide_texts)
            if slide_text.strip():
                slide_chunks = _chunk_text(slide_text, chunk_size, chunk_overlap, page_number=slide_num)
                chunks.extend(slide_chunks)
        return chunks

    # Fallback
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        full_text = f.read()
    return _chunk_text(full_text, chunk_size, chunk_overlap)


def _chunk_text(text: str, chunk_size: int = 800, chunk_overlap: int = 150, page_number: int = 1) -> List[Dict[str, Any]]:
    chunks = []
    cleaned_text = " ".join(text.split())
    if not cleaned_text:
        return []

    start = 0
    idx = 0
    while start < len(cleaned_text):
        end = min(start + chunk_size, len(cleaned_text))
        chunk_content = cleaned_text[start:end]
        chunks.append({
            "chunk_index": idx,
            "content": chunk_content,
            "page_number": page_number,
            "metadata_json": {
                "char_length": len(chunk_content),
                "page": page_number
            }
        })
        idx += 1
        if end == len(cleaned_text):
            break
        start += (chunk_size - chunk_overlap)

    return chunks
